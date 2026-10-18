from datetime import date, datetime
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook

from storefront.models import ItineraryReference, ItineraryStop


def _text(value):
    return "" if value is None else str(value).strip()


def _date_for_stop(value, year):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    value = _text(value)
    if not value:
        return None
    try:
        return datetime.strptime(f"{value}/{year}", "%m/%d/%Y").date()
    except ValueError as exc:
        raise CommandError(f"Invalid itinerary date: {value}") from exc


class Command(BaseCommand):
    help = "Import itinerary stops, planning assumptions, and sources from the RV planner workbook."

    def add_arguments(self, parser):
        parser.add_argument(
            "--source",
            default=str(settings.BASE_DIR / "2027_RV_Master_Trip_Planner.xlsx"),
            help="Path to the planner workbook (defaults to the project root).",
        )
        parser.add_argument(
            "--update-existing",
            action="store_true",
            help="Refresh existing workbook fields without replacing manually entered coordinates.",
        )

    def handle(self, *args, **options):
        workbook_path = Path(options["source"])
        if not workbook_path.is_file():
            raise CommandError(f"Workbook not found: {workbook_path}")

        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        if "2027 Master Itinerary" not in workbook.sheetnames:
            raise CommandError("Workbook is missing the '2027 Master Itinerary' sheet.")

        created = updated = skipped = 0
        for row in workbook["2027 Master Itinerary"].iter_rows(values_only=True):
            if len(row) < 20 or row[0] is None or not row[1] or not _text(row[2]):
                continue
            stop_date = _date_for_stop(row[1], date.today().year)
            check_in = _date_for_stop(row[5], stop_date.year)
            check_out = _date_for_stop(row[6], stop_date.year)
            if check_in and check_out and check_out < check_in:
                check_out = check_out.replace(year=check_out.year + 1)

            values = {
                "day_number": int(row[0]),
                "stop_date": stop_date,
                "destination": _text(row[2]),
                "preferred_campground": _text(row[3]),
                "alternate_campground": _text(row[4]),
                "check_in": check_in,
                "check_out": check_out,
                "route": _text(row[7]),
                "drive_time": _text(row[8]),
                "site_type": _text(row[9]),
                "booking_requirements": _text(row[10]),
                "price_or_membership": _text(row[11]),
                "booking_opening_estimate": _text(row[12]),
                "attractions": _text(row[13]),
                "landscape": _text(row[14]),
                "fuel_errands": _text(row[15]),
                "stop_category": _text(row[16]),
                "weather": _text(row[17]),
                "towing_notes": _text(row[18]),
                "booking_status": (
                    ItineraryStop.BOOKED
                    if _text(row[19]).lower() in ("booked", "final")
                    else ItineraryStop.BOOK
                ),
            }
            if len(row) > 20:
                values["booking_confirmation"] = _text(row[20])
            stop, was_created = ItineraryStop.objects.get_or_create(
                stop_date=stop_date,
                destination=values["destination"],
                defaults=values,
            )
            if was_created:
                created += 1
            elif options["update_existing"]:
                for field, value in values.items():
                    setattr(stop, field, value)
                stop.save()
                updated += 1
            else:
                skipped += 1

        self._import_references(workbook)
        workbook.close()
        self.stdout.write(
            self.style.SUCCESS(f"Imported {created} stops; updated {updated}; skipped {skipped} existing stops.")
        )

    def _import_references(self, workbook):
        if "Planning Notes" in workbook.sheetnames:
            for order, row in enumerate(workbook["Planning Notes"].iter_rows(values_only=True), start=1):
                if not row or not _text(row[0]) or not _text(row[1] if len(row) > 1 else None):
                    continue
                ItineraryReference.objects.get_or_create(
                    reference_type=ItineraryReference.ASSUMPTION,
                    title=_text(row[0]),
                    defaults={"detail": _text(row[1]), "display_order": order},
                )

        if "Sources" in workbook.sheetnames:
            for order, row in enumerate(workbook["Sources"].iter_rows(min_row=2, values_only=True), start=1):
                if len(row) < 3 or not _text(row[0]):
                    continue
                ItineraryReference.objects.get_or_create(
                    reference_type=ItineraryReference.SOURCE,
                    title=_text(row[0]),
                    defaults={"url": _text(row[1]), "detail": _text(row[2]), "display_order": order},
                )