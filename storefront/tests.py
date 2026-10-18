import tempfile
from datetime import date, datetime
from io import StringIO

from django.contrib.auth.models import Group, User
from django.core.management import call_command
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook

from config.hashers import LegacySHA256PasswordHasher
from .models import ItineraryReference, ItineraryStop, Organization, TripEntry, TripEntryFile


class TripSiteTests(TestCase):
    def setUp(self):
        temporary_media = tempfile.TemporaryDirectory()
        self.addCleanup(temporary_media.cleanup)
        media_settings = override_settings(MEDIA_ROOT=temporary_media.name)
        media_settings.enable()
        self.addCleanup(media_settings.disable)

        self.admin = User.objects.create_user("admin", password="FarmRoad!2026")
        self.admin.groups.add(Group.objects.get(name="ADMIN"))
        self.viewer = User.objects.create_user("viewer", password="FarmRoad!2026")
        self.viewer.groups.add(Group.objects.get(name="USER"))

    def test_home_page_shows_trip_sign_in_not_egg_placeholder(self):
        response = self.client.get(reverse("home"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Make room for")
        self.assertContains(response, "MEMBER SIGN IN")
        self.assertNotContains(response, "Fresh from our little corner")

    def test_legacy_sha256_login_upgrades_password_hash(self):
        user = User.objects.create(username="legacy-user")
        user.password = LegacySHA256PasswordHasher().encode("old-trip-password", "")
        user.save(update_fields=["password"])

        response = self.client.post(
            reverse("home"),
            {"username": "legacy-user", "password": "old-trip-password"},
        )

        self.assertRedirects(response, reverse("home"))
        user.refresh_from_db()
        self.assertTrue(user.password.startswith("pbkdf2_sha256$"))
        self.assertContains(self.client.get(reverse("home")), "Hello,")

    def test_trip_journal_requires_sign_in(self):
        response = self.client.get(reverse("trip_index"))

        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.url.startswith("/?next="))

    def test_user_can_create_and_download_private_trip_file(self):
        self.client.force_login(self.viewer)
        response = self.client.post(
            reverse("create_trip_entry"),
            {
                "event_title": "Arrive in London",
                "event_date": "2027-06-10",
                "event_description": "First day of the trip.",
                "file_descriptions": ["Train ticket"],
                "files": [SimpleUploadedFile("ticket.txt", b"ticket contents", "text/plain")],
            },
        )

        entry = TripEntry.objects.get(event_title="Arrive in London")
        trip_file = TripEntryFile.objects.get(trip_entry=entry)
        self.assertRedirects(response, reverse("trip_entry", args=[entry.pk]))
        self.assertEqual(trip_file.description, "Train ticket")
        self.assertTrue(trip_file.file.name.startswith("trip_2027_files/other/"))

        self.client.logout()
        denied = self.client.get(reverse("download_trip_file", args=[trip_file.pk]))
        self.assertEqual(denied.status_code, 302)

        self.client.force_login(self.viewer)
        download = self.client.get(reverse("download_trip_file", args=[trip_file.pk]))
        self.assertEqual(download.status_code, 200)
        self.assertEqual(b"".join(download.streaming_content), b"ticket contents")
        self.assertEqual(download["X-Content-Type-Options"], "nosniff")

    def test_trip_entry_coordinates_are_validated_and_crud_supported(self):
        self.client.force_login(self.viewer)
        created = self.client.post(
            reverse("create_trip_entry"),
            {
                "event_title": "Coordinates",
                "event_date": "2027-06-10",
                "event_description": "Location test",
                "geo_location": " 51.5072, -0.1276 ",
            },
        )
        entry = TripEntry.objects.get(event_title="Coordinates")
        self.assertRedirects(created, reverse("trip_entry", args=[entry.pk]))
        self.assertEqual(entry.geo_location, "51.5072, -0.1276")

        updated = self.client.post(
            reverse("trip_entry", args=[entry.pk]),
            {
                "event_title": "Coordinates updated",
                "event_date": "2027-06-11",
                "event_description": "Updated location",
                "geo_location": "48.8566, 2.3522",
            },
        )
        self.assertRedirects(updated, reverse("trip_entry", args=[entry.pk]))
        entry.refresh_from_db()
        self.assertEqual(entry.geo_location, "48.8566, 2.3522")

        invalid = self.client.post(
            reverse("trip_entry", args=[entry.pk]),
            {
                "event_title": "Invalid coordinates",
                "event_date": "2027-06-11",
                "event_description": "Invalid location",
                "geo_location": "91, 181",
            },
        )
        self.assertEqual(invalid.status_code, 400)
        entry.refresh_from_db()
        self.assertEqual(entry.geo_location, "48.8566, 2.3522")

    def test_trip_map_payload_orders_located_events_by_date(self):
        later = TripEntry.objects.create(
            event_date=date(2027, 6, 12),
            event_title="Later stop",
            geo_location="48.8566, 2.3522",
        )
        earlier = TripEntry.objects.create(
            event_date=date(2027, 6, 10),
            event_title="First stop",
            geo_location="51.5072, -0.1276",
        )
        TripEntry.objects.create(
            event_date=date(2027, 6, 11),
            event_title="No coordinates",
        )
        self.client.force_login(self.viewer)

        response = self.client.get(reverse("trip_index"))

        self.assertContains(response, 'data-trip-map')
        self.assertContains(response, "maplibre-gl")
        self.assertContains(response, "/static/storefront/map.js")
        self.assertNotContains(response, "GOOGLE_MAPS_API_KEY")
        first_position = response.content.index(b'"title": "First stop"')
        later_position = response.content.index(b'"title": "Later stop"')
        self.assertLess(first_position, later_position)
        self.assertNotContains(response, '"title": "No coordinates"')
        self.assertEqual(earlier.geo_location, "51.5072, -0.1276")
        self.assertEqual(later.geo_location, "48.8566, 2.3522")

    def test_mp4_playback_supports_private_byte_ranges(self):
        entry = TripEntry.objects.create(event_date=date(2027, 6, 10), event_title="Video playback")
        trip_file = TripEntryFile.objects.create(
            trip_entry=entry,
            file=SimpleUploadedFile("clip.mp4", b"0123456789", "video/mp4"),
            file_name="clip.mp4",
            file_type="video/mp4",
        )
        self.client.force_login(self.viewer)

        partial = self.client.get(reverse("download_trip_file", args=[trip_file.pk]), HTTP_RANGE="bytes=2-5")

        self.assertEqual(partial.status_code, 206)
        self.assertEqual(b"".join(partial.streaming_content), b"2345")
        self.assertEqual(partial["Content-Range"], "bytes 2-5/10")
        self.assertEqual(partial["Content-Length"], "4")
        self.assertEqual(partial["Content-Type"], "video/mp4")
        self.assertEqual(partial["Accept-Ranges"], "bytes")
        self.assertTrue(partial["Content-Disposition"].startswith("inline;"))

        suffix = self.client.get(reverse("download_trip_file", args=[trip_file.pk]), HTTP_RANGE="bytes=-3")
        self.assertEqual(suffix.status_code, 206)
        self.assertEqual(b"".join(suffix.streaming_content), b"789")

        unsatisfiable = self.client.get(reverse("download_trip_file", args=[trip_file.pk]), HTTP_RANGE="bytes=10-")
        self.assertEqual(unsatisfiable.status_code, 416)
        self.assertEqual(unsatisfiable["Content-Range"], "bytes */10")

        self.client.logout()
        denied = self.client.get(reverse("download_trip_file", args=[trip_file.pk]), HTTP_RANGE="bytes=0-1")
        self.assertEqual(denied.status_code, 302)

    def test_chunk_upload_succeeds_with_csrf_enforcement(self):
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.viewer)
        page = client.get(reverse("trip_index"))
        csrf_token = page.cookies["csrftoken"].value
        create_response = client.post(
            reverse("create_trip_entry"),
            {
                "event_title": "CSRF upload",
                "event_date": "2027-06-10",
                "csrfmiddlewaretoken": csrf_token,
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
            HTTP_X_CSRFTOKEN=csrf_token,
        )
        self.assertEqual(create_response.status_code, 200)

        upload_response = client.post(
            create_response.json()["upload_url"],
            {
                "upload_id": "b1b2c3d4-e5f6-47a8-9012-1234567890ab",
                "offset": "0",
                "total_size": "4",
                "file_name": "csrf.mp4",
                "file_type": "video/mp4",
                "description": "",
                "csrfmiddlewaretoken": csrf_token,
                "chunk": SimpleUploadedFile("csrf.mp4", b"data", "video/mp4"),
            },
        )
        self.assertEqual(upload_response.status_code, 200, upload_response.content)
        self.assertTrue(upload_response.json()["complete"])
        self.assertTrue(TripEntryFile.objects.filter(file_name="csrf.mp4").exists())

    def test_non_mp4_video_is_rejected_from_event_form(self):
        self.client.force_login(self.viewer)

        response = self.client.post(
            reverse("create_trip_entry"),
            {
                "event_title": "Unsupported video",
                "event_date": "2027-06-10",
                "files": [SimpleUploadedFile("clip.avi", b"video-data", "video/x-msvideo")],
            },
        )

        self.assertRedirects(response, "/trip-2027/?error=video-format-not-allowed")
        self.assertFalse(TripEntry.objects.filter(event_title="Unsupported video").exists())

    def test_non_mp4_video_is_rejected_from_chunk_upload(self):
        self.client.force_login(self.viewer)
        entry = TripEntry.objects.create(event_date=date(2027, 6, 10), event_title="Unsupported chunk")

        response = self.client.post(
            reverse("upload_trip_file", args=[entry.pk]),
            {
                "upload_id": "c1b2c3d4-e5f6-47a8-9012-1234567890ab",
                "offset": "0",
                "total_size": "4",
                "file_name": "clip.mpg",
                "file_type": "application/octet-stream",
                "description": "",
                "chunk": SimpleUploadedFile("clip.mpg", b"data", "application/octet-stream"),
            },
        )

        self.assertEqual(response.status_code, 415)
        self.assertEqual(response.json()["error"], "video-format-not-allowed")
        self.assertFalse(entry.files.exists())

    def test_user_can_upload_mp4_larger_than_previous_limit_in_chunks(self):
        self.client.force_login(self.viewer)
        response = self.client.post(
            reverse("create_trip_entry"),
            {
                "event_title": "Long video",
                "event_date": "2027-06-10",
                "event_description": "A large video upload.",
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(response.status_code, 200)
        upload_data = response.json()
        total_size = 20 * 1024 * 1024 + 1
        chunk_size = 5 * 1024 * 1024
        upload_id = "a1b2c3d4-e5f6-47a8-9012-1234567890ab"

        for offset in range(0, total_size, chunk_size):
            chunk = b"v" * min(chunk_size, total_size - offset)
            chunk_response = self.client.post(
                upload_data["upload_url"],
                {
                    "upload_id": upload_id,
                    "offset": str(offset),
                    "total_size": str(total_size),
                    "file_name": "journey.mp4",
                    "file_type": "video/mp4",
                    "description": "Road footage",
                    "chunk": SimpleUploadedFile("journey.mp4", chunk, "application/octet-stream"),
                },
            )
            self.assertEqual(chunk_response.status_code, 200, chunk_response.content)

        trip_file = TripEntryFile.objects.get(file_name="journey.mp4")
        self.assertEqual(trip_file.file.size, total_size)
        self.assertEqual(trip_file.file_type, "video/mp4")
        self.assertEqual(trip_file.description, "Road footage")
        self.assertTrue(trip_file.file.name.startswith("trip_2027_files/videos/"))

    def test_admin_can_manage_organizations_and_users_cannot(self):
        self.client.force_login(self.viewer)
        denied = self.client.get(reverse("admin_users"))
        self.assertEqual(denied.status_code, 302)
        self.assertEqual(self.client.get(reverse("admin_data_dictionary")).status_code, 302)
        self.assertNotContains(self.client.get(reverse("trip_index")), "Data Dictionary")

        self.client.force_login(self.admin)
        response = self.client.post(
            reverse("admin_organizations"),
            {"action": "create", "business_name": "Joe's Trip", "city": "London"},
        )
        self.assertRedirects(response, reverse("admin_organizations"))
        self.assertTrue(Organization.objects.filter(business_name="Joe's Trip", city="London").exists())
        self.assertEqual(self.client.get(reverse("admin_roles")).status_code, 200)
        self.assertContains(self.client.get(reverse("admin_itinerary")), "Itinerary Planner")
        dictionary_response = self.client.get(reverse("admin_data_dictionary"))
        self.assertContains(dictionary_response, "Data Dictionary")
        self.assertContains(dictionary_response, "storefront_tripentryfile")
        self.assertContains(dictionary_response, "trip_entry_id")

    def test_eggs_route_is_role_restricted(self):
        self.client.force_login(self.viewer)
        self.assertEqual(self.client.get(reverse("eggs")).status_code, 302)

        eggs_user = User.objects.create_user("eggs", password="FarmRoad!2026")
        eggs_user.groups.add(Group.objects.get(name="EGGS"))
        self.client.force_login(eggs_user)
        self.assertContains(self.client.get(reverse("eggs")), "EGGS area")

    def test_authenticated_journal_and_admin_pages_render(self):
        entry = TripEntry.objects.create(
            event_date=date(2027, 6, 10),
            event_title="First day",
            event_description="A new city.",
            geo_location="51.5072, -0.1276",
        )
        TripEntryFile.objects.create(
            trip_entry=entry,
            file=SimpleUploadedFile("arrival.jpg", b"image-bytes", "image/jpeg"),
            file_name="arrival.jpg",
            file_type="image/jpeg",
            description="First morning",
        )
        self.assertTrue(entry.files.get().file.name.startswith("trip_2027_files/images/"))
        TripEntryFile.objects.create(
            trip_entry=entry,
            file=SimpleUploadedFile("clip.mp4", b"video-bytes", "video/mp4"),
            file_name="clip.mp4",
            file_type="video/mp4",
            description="A short clip",
        )
        Organization.objects.create(business_name="Trip office", city="London")
        self.client.force_login(self.admin)

        for route in [
            "trip_index",
            "admin_users",
            "admin_roles",
            "admin_organizations",
            "admin_data_dictionary",
            "admin_itinerary",
        ]:
            with self.subTest(route=route):
                self.assertEqual(self.client.get(reverse(route)).status_code, 200)
        journal = self.client.get(reverse("trip_index"))
        self.assertContains(journal, 'data-map-disclosure open')
        self.assertContains(journal, 'data-trip-map')
        detail = self.client.get(reverse("trip_entry", args=[entry.pk]))
        self.assertContains(detail, "First day")
        self.assertContains(detail, "arrival.jpg")
        self.assertContains(detail, "First morning")
        self.assertContains(detail, "data-video-viewer")
        self.assertContains(detail, "data-video-open")
        self.assertContains(detail, "controls")
        self.assertContains(detail, "data-upload-overall")

    def test_trip_map_includes_itinerary_overlay_for_all_viewers(self):
        ItineraryStop.objects.create(
            day_number=1,
            stop_date=date(2027, 5, 1),
            destination="Seaside, OR",
            geo_location="45.9932, -123.9226",
        )

        self.client.force_login(self.viewer)
        viewer_response = self.client.get(reverse("trip_index"))
        self.assertContains(viewer_response, 'data-itinerary-toggle checked')
        self.assertContains(viewer_response, "45.9932")
        viewer_itinerary_events = viewer_response.context["itinerary_map_events"]
        self.assertEqual(len(viewer_itinerary_events), 1)
        self.assertNotIn("url", viewer_itinerary_events[0])

        self.client.force_login(self.admin)
        admin_response = self.client.get(reverse("trip_index"))
        self.assertContains(admin_response, 'data-itinerary-toggle checked')
        self.assertContains(admin_response, 'data-map-mode="trip"')
        self.assertEqual(
            admin_response.context["itinerary_map_events"],
            [
                {
                    "lat": 45.9932,
                    "lng": -123.9226,
                    "title": "Seaside, OR",
                    "date": "2027-05-01",
                    "url": f"{reverse('admin_itinerary')}#stop-1",
                }
            ],
        )

    def test_itinerary_planner_is_admin_only_and_validates_coordinates(self):
        self.client.force_login(self.viewer)
        denied = self.client.get(reverse("admin_itinerary"))
        self.assertEqual(denied.status_code, 302)
        self.assertNotContains(self.client.get(reverse("trip_index")), "Itinerary Planner")

        self.client.force_login(self.admin)
        self.assertContains(self.client.get(reverse("trip_index")), reverse("admin_itinerary"))
        page = self.client.get(reverse("admin_itinerary"))
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, "Add a destination")
        self.assertContains(page, "Booking confirmation")

        created = self.client.post(
            reverse("admin_itinerary"),
            {
                "action": "create",
                "day_number": 1,
                "stop_date": "2027-05-01",
                "destination": "Seaside, OR",
                "preferred_campground": "Fort Stevens State Park",
                "booking_confirmation": "CONF-123",
                "booking_status": "Book",
                "geo_location": " 45.9932, -123.9226 ",
            },
        )
        self.assertRedirects(created, reverse("admin_itinerary"))
        stop = ItineraryStop.objects.get(destination="Seaside, OR")
        self.assertEqual(stop.geo_location, "45.9932, -123.9226")
        self.assertEqual(stop.booking_confirmation, "CONF-123")
        self.assertContains(self.client.get(reverse("admin_itinerary")), "data-trip-map")
        self.assertContains(self.client.get(reverse("admin_itinerary")), 'data-map-disclosure open')

        invalid = self.client.post(
            reverse("admin_itinerary"),
            {
                "action": "create",
                "day_number": 2,
                "stop_date": "2027-05-03",
                "destination": "Invalid coordinates",
                "geo_location": "91, -123.9226",
            },
        )
        self.assertEqual(invalid.status_code, 400)
        self.assertFalse(ItineraryStop.objects.filter(destination="Invalid coordinates").exists())

    def test_itinerary_date_change_and_delete_shift_unbooked_stops(self):
        self.client.force_login(self.admin)
        first = ItineraryStop.objects.create(day_number=1, stop_date=date(2027, 5, 1), destination="A")
        unbooked = ItineraryStop.objects.create(day_number=2, stop_date=date(2027, 5, 3), destination="B")
        booked = ItineraryStop.objects.create(
            day_number=3, stop_date=date(2027, 5, 5), destination="C", booking_status="Booked"
        )
        self.client.post(
            reverse("admin_itinerary"),
            {
                "action": "update", "form_type": "collision", "stop_id": first.pk,
                "stop_date": "2027-05-02", "destination": "A", "booking_status": "Book",
            },
        )
        unbooked.refresh_from_db()
        booked.refresh_from_db()
        self.assertEqual(unbooked.stop_date, date(2027, 5, 4))
        self.assertEqual(booked.stop_date, date(2027, 5, 5))

        response = self.client.get(reverse("admin_itinerary"))
        self.assertNotContains(response, "Collisions")
        self.client.post(reverse("admin_itinerary"), {"action": "delete", "stop_id": first.pk})
        unbooked.refresh_from_db()
        self.assertEqual(unbooked.stop_date, date(2027, 5, 3))
        booked.stop_date = date(2027, 5, 3)
        booked.save()
        self.assertContains(self.client.get(reverse("admin_itinerary")), "Collisions")

    def test_itinerary_workbook_import_is_repeatable_and_preserves_coordinates(self):
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = f"{directory}/planner.xlsx"
            workbook = Workbook()
            itinerary = workbook.active
            itinerary.title = "2027 Master Itinerary"
            itinerary.append(
                [
                    1, datetime(2027, 5, 1), "Seaside, OR", "Thousand Trails Seaside",
                    "Fort Stevens State Park", "5/1", "5/3", "Portland -> Seaside", "~1.5 hr",
                    "RV site", "Confirm at booking", "TT member", "~Mar 2", "Cannon Beach",
                    "Beach", "Fuel in Astoria", "Campground", "Cool/mild", "Easy tow", "BOOK",
                    "RES-2027-001",
                ]
            )
            notes = workbook.create_sheet("Planning Notes")
            notes.append(["Planning assumptions", None])
            notes.append(["Daily towing limit", "Target maximum 5 hours."])
            sources = workbook.create_sheet("Sources")
            sources.append(["Source", "URL", "Use"])
            sources.append(["Reservation guide", "https://example.com/reservations", "Booking rules"])
            workbook.save(workbook_path)

            call_command("import_itinerary_workbook", source=workbook_path, stdout=StringIO())
            stop = ItineraryStop.objects.get(destination="Seaside, OR")
            self.assertEqual(stop.stop_date, date(2027, 5, 1))
            self.assertEqual(stop.check_out, date(2027, 5, 3))
            self.assertEqual(stop.price_or_membership, "TT member")
            self.assertEqual(stop.booking_confirmation, "RES-2027-001")
            self.assertEqual(ItineraryReference.objects.count(), 2)

            stop.geo_location = "45.9932, -123.9226"
            stop.save(update_fields=["geo_location"])
            call_command("import_itinerary_workbook", source=workbook_path, stdout=StringIO())
            self.assertEqual(ItineraryStop.objects.count(), 1)
            self.assertEqual(ItineraryStop.objects.get().geo_location, "45.9932, -123.9226")

    def test_admin_can_create_user_and_role(self):
        self.client.force_login(self.admin)
        user_response = self.client.post(
            reverse("admin_users"),
            {
                "action": "create",
                "username": "new-member",
                "first_name": "New",
                "last_name": "Member",
                "password": "DifferentSecure!2026",
                "roles": [Group.objects.get(name="USER").pk],
            },
        )
        self.assertRedirects(user_response, reverse("admin_users"))
        self.assertTrue(User.objects.get(username="new-member").groups.filter(name="USER").exists())

        role_response = self.client.post(
            reverse("admin_roles"),
            {"action": "create", "name": "TRAVEL"},
        )
        self.assertRedirects(role_response, reverse("admin_roles"))
        self.assertTrue(Group.objects.filter(name="TRAVEL").exists())

    def test_user_can_delete_trip_file_from_local_media(self):
        entry = TripEntry.objects.create(
            event_date=date(2027, 6, 10),
            event_title="Delete file test",
        )
        trip_file = TripEntryFile.objects.create(
            trip_entry=entry,
            file=SimpleUploadedFile("delete-me.txt", b"temporary file", "text/plain"),
            file_name="delete-me.txt",
            file_type="text/plain",
        )
        stored_name = trip_file.file.name
        self.client.force_login(self.viewer)

        response = self.client.post(
            reverse("delete_trip_file", args=[entry.pk, trip_file.pk]),
        )

        self.assertRedirects(response, reverse("trip_entry", args=[entry.pk]))
        self.assertFalse(TripEntryFile.objects.filter(pk=trip_file.pk).exists())
        self.assertFalse(trip_file.file.storage.exists(stored_name))

    def test_admin_can_delete_trip_entry_and_attached_files(self):
        entry = TripEntry.objects.create(
            event_date=date(2027, 6, 11),
            event_title="Delete event test",
        )
        trip_file = TripEntryFile.objects.create(
            trip_entry=entry,
            file=SimpleUploadedFile("delete-event-file.txt", b"temporary file", "text/plain"),
            file_name="delete-event-file.txt",
            file_type="text/plain",
        )
        stored_name = trip_file.file.name
        self.client.force_login(self.admin)

        response = self.client.post(reverse("delete_trip_entry", args=[entry.pk]))

        self.assertRedirects(response, reverse("trip_index"))
        self.assertFalse(TripEntry.objects.filter(pk=entry.pk).exists())
        self.assertFalse(TripEntryFile.objects.filter(pk=trip_file.pk).exists())
        self.assertFalse(trip_file.file.storage.exists(stored_name))
